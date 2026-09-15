#!/usr/bin/env python3
"""
Convertit un fichier prospects_*.xlsx (généré par prospection_osm.py) en CSV
prêt à importer dans Brevo : ne garde que les lignes avec un email, et
formate les colonnes proprement.

Installation :
    pip install openpyxl --break-system-packages

Usage :
    python3 prep_csv_brevo.py prospects_boucherie_halal_Lyon.xlsx
    python3 prep_csv_brevo.py prospects_toutes_categories_Lyon.xlsx --ville Lyon

Résultat : un fichier contacts_brevo.csv dans le dossier courant.
"""

import argparse
import csv
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("openpyxl manquant : pip install openpyxl --break-system-packages")
    sys.exit(1)


def deviner_ville(chemin_fichier: str) -> str:
    """Essaie d'extraire la ville depuis le nom du fichier prospects_xxx_Ville.xlsx"""
    nom = Path(chemin_fichier).stem
    parts = nom.split("_")
    return parts[-1] if parts else ""


def main():
    parser = argparse.ArgumentParser(description="Prépare un CSV Brevo depuis un export prospection_osm.py")
    parser.add_argument("fichier_xlsx", help="Chemin vers le fichier prospects_*.xlsx")
    parser.add_argument("--ville", help="Nom de la ville (sinon déduit du nom de fichier)")
    parser.add_argument("--sortie", default="contacts_brevo.csv", help="Nom du CSV de sortie")
    args = parser.parse_args()

    chemin = Path(args.fichier_xlsx)
    if not chemin.exists():
        print(f"Fichier introuvable : {chemin}")
        sys.exit(1)

    ville = args.ville or deviner_ville(args.fichier_xlsx)

    wb = load_workbook(chemin, data_only=True)

    # Si le fichier a un onglet "TOUS" (généré avec --toutes), on l'utilise.
    # Sinon on prend le premier (et seul) onglet.
    if "TOUS" in wb.sheetnames:
        ws = wb["TOUS"]
        a_colonne_categorie = True
    else:
        ws = wb[wb.sheetnames[0]]
        a_colonne_categorie = False
        categorie_unique = wb.sheetnames[0]

    lignes = list(ws.iter_rows(values_only=True))
    if not lignes:
        print("Feuille vide.")
        sys.exit(0)

    entetes = [str(h).strip().lower() if h else "" for h in lignes[0]]

    def idx(nom_colonne):
        return entetes.index(nom_colonne) if nom_colonne in entetes else None

    i_nom = idx("nom")
    i_adresse = idx("adresse")
    i_tel = idx("telephone")
    i_email = idx("email")
    i_site = idx("site_web")
    i_categorie = idx("categorie")

    contacts = []
    for row in lignes[1:]:
        email = (row[i_email] or "").strip() if i_email is not None else ""
        if not email or "@" not in email:
            continue  # Brevo n'a pas besoin des lignes sans email exploitable

        contacts.append({
            "EMAIL": email,
            "NOM_ENTREPRISE": row[i_nom] or "" if i_nom is not None else "",
            "VILLE": ville,
            "CATEGORIE": (row[i_categorie] if a_colonne_categorie and i_categorie is not None else categorie_unique) or "",
            "TELEPHONE": row[i_tel] or "" if i_tel is not None else "",
            "ADRESSE": row[i_adresse] or "" if i_adresse is not None else "",
            "SITE_WEB": row[i_site] or "" if i_site is not None else "",
        })

    if not contacts:
        print("Aucun contact avec email trouvé dans ce fichier.")
        print("Astuce : relance prospection_osm.py sans --no-enrichir pour aller chercher les emails sur les sites.")
        sys.exit(0)

    # dédoublonnage par email (au cas où)
    vus = set()
    contacts_uniques = []
    for c in contacts:
        if c["EMAIL"].lower() not in vus:
            vus.add(c["EMAIL"].lower())
            contacts_uniques.append(c)

    colonnes = ["EMAIL", "NOM_ENTREPRISE", "VILLE", "CATEGORIE", "TELEPHONE", "ADRESSE", "SITE_WEB"]
    with open(args.sortie, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=colonnes)
        writer.writeheader()
        writer.writerows(contacts_uniques)

    print(f"{len(contacts_uniques)} contacts avec email écrits dans {args.sortie}")
    print("Colonnes prêtes pour l'import Brevo : EMAIL, NOM_ENTREPRISE, VILLE, CATEGORIE, TELEPHONE, ADRESSE, SITE_WEB")


if __name__ == "__main__":
    main()
