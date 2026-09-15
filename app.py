#!/usr/bin/env python3
"""
Prospection automatisée : trouve les entreprises d'une ville, avec ou sans
site web, et récupère leur téléphone/email (OSM + scraping léger du site
si besoin), via OpenStreetMap.

Gratuit, pas de clé API. Respecte les usage policies de Nominatim/Overpass
(1 req/s max, User-Agent obligatoire) -> ne pas paralléliser.

Installation :
    pip install requests openpyxl --break-system-packages

Usage :
    python3 app.py --categorie plombier --ville "Troyes, France"
    python3 app.py --categorie coiffeur --ville "Reims, France" --rayon 8000
    python3 app.py --toutes --ville "Lyon, France"
    python3 app.py --liste-categories
    python3 app.py --categorie boucherie_halal --ville "Paris, France" --no-enrichir

Résultat : un fichier .xlsx dans le dossier courant. Avec --toutes, un
onglet par catégorie + un onglet "TOUS" trié par score de priorité.
"""

import argparse
import re
import sys
import time
import requests

HEADERS = {"User-Agent": "ProspectionFreelance/1.0 (contact: ton-email@exemple.fr)"}

NOMINATIM_URL = "https://nominatim.openstreetmap.org/search"
OVERPASS_URL = "https://overpass-api.de/api/interpreter"

CATEGORIES = {
    "plombier": ['["craft"="plumber"]'],
    "electricien": ['["craft"="electrician"]'],
    "coiffeur": ['["shop"="hairdresser"]'],
    "restaurant": ['["amenity"="restaurant"]'],
    "boulangerie": ['["shop"="bakery"]'],
    "avocat": ['["office"="lawyer"]'],
    "architecte": ['["office"="architect"]'],
    "comptable": ['["office"="accountant"]'],
    "garage": ['["shop"="car_repair"]'],
    "fleuriste": ['["shop"="florist"]'],
    "dentiste": ['["amenity"="dentist"]'],
    "kine": ['["healthcare"="physiotherapist"]'],
    "immobilier": ['["office"="estate_agent"]'],
    "menuisier": ['["craft"="carpenter"]'],
    "peintre_batiment": ['["craft"="painter"]'],
    "boucherie": ['["shop"="butcher"]'],
    "opticien": ['["shop"="optician"]'],
    "veterinaire": ['["amenity"="veterinary"]'],
    "traiteur": ['["shop"="catering"]'],
    "photographe": ['["craft"="photographer"]'],
    "boucherie_halal": ['["shop"="butcher"]["diet:halal"="yes"]'],
    "epicerie_halal": ['["shop"~"grocery|convenience|supermarket"]["diet:halal"="yes"]'],
    "traiteur_halal": ['["shop"="catering"]["diet:halal"="yes"]', '["shop"="deli"]["diet:halal"="yes"]'],
    "patisserie_halal": ['["shop"="pastry"]["diet:halal"="yes"]'],
    "boulangerie_halal": ['["shop"="bakery"]["diet:halal"="yes"]'],
    "fast_food_halal": ['["amenity"="fast_food"]["diet:halal"="yes"]'],
}

EMAIL_RE = re.compile(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]{2,}')
PHONE_RE = re.compile(r'0[1-9](?:[ .\-]?\d{2}){4}')
EMAIL_DOMAINES_A_IGNORER = (
    "sentry.io", "wixpress.com", "example.com", "domain.com",
    "godaddy.com", "yourdomain.com", ".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp",
)


def geocode_ville(ville: str):
    params = {"q": ville, "format": "json", "limit": 1}
    r = requests.get(NOMINATIM_URL, params=params, headers=HEADERS, timeout=15)
    r.raise_for_status()
    data = r.json()
    if not data:
        raise ValueError(f"Ville introuvable : {ville}")
    return float(data[0]["lat"]), float(data[0]["lon"]), data[0]["display_name"]


def query_overpass(tag_filters, lat, lon, rayon_m):
    filtres = "".join(
        f'node{tf}(around:{rayon_m},{lat},{lon});way{tf}(around:{rayon_m},{lat},{lon});'
        for tf in tag_filters
    )
    query = f"[out:json][timeout:60];({filtres});out center tags;"
    r = requests.post(OVERPASS_URL, data={"data": query}, headers=HEADERS, timeout=90)
    r.raise_for_status()
    return r.json().get("elements", [])


def extraire_adresse(tags: dict) -> str:
    rue = " ".join(p for p in [tags.get("addr:housenumber", ""), tags.get("addr:street", "")] if p).strip()
    ville = tags.get("addr:city", "")
    cp = tags.get("addr:postcode", "")
    return ", ".join(p for p in [rue, f"{cp} {ville}".strip()] if p)


def a_un_site(tags: dict) -> str:
    site = tags.get("website") or tags.get("contact:website") or ""
    if site and not site.startswith("http"):
        site = "https://" + site
    return site


def chercher_contact_sur_site(url: str):
    """Va chercher un email et un téléphone sur la page d'accueil + page contact d'un site."""
    email_trouve, tel_trouve = "", ""
    pages_a_tester = [url]

    try:
        r = requests.get(url, headers=HEADERS, timeout=8)
        html = r.text
        # tente de repérer un lien "contact" dans la page d'accueil
        m = re.search(r'href=["\']([^"\']*contact[^"\']*)["\']', html, re.IGNORECASE)
        if m:
            lien = m.group(1)
            if lien.startswith("/"):
                from urllib.parse import urljoin
                lien = urljoin(url, lien)
            if lien.startswith("http"):
                pages_a_tester.append(lien)
    except requests.RequestException:
        return "", ""

    for page in pages_a_tester[:2]:
        try:
            r = requests.get(page, headers=HEADERS, timeout=8)
            html = r.text
        except requests.RequestException:
            continue

        if not email_trouve:
            for m in EMAIL_RE.findall(html):
                if not any(d in m.lower() for d in EMAIL_DOMAINES_A_IGNORER):
                    email_trouve = m
                    break

        if not tel_trouve:
            m = PHONE_RE.search(html)
            if m:
                tel_trouve = m.group(0)

        if email_trouve and tel_trouve:
            break

        time.sleep(0.4)  # politesse envers le site tiers

    return email_trouve, tel_trouve


def score_priorite(a_site: bool, tel: str, email: str, nom: str) -> int:
    score = 0
    if not a_site:
        score += 3
    if tel:
        score += 1
    if email:
        score += 1
    if nom:
        score += 1
    return score


def collecter_categorie(categorie, lat, lon, rayon, enrichir):
    print(f"Recherche des '{categorie}' dans un rayon de {rayon}m...")
    elements = query_overpass(CATEGORIES[categorie], lat, lon, rayon)
    print(f"  -> {len(elements)} résultats bruts")

    resultats = []
    for el in elements:
        tags = el.get("tags", {})
        nom = tags.get("name")
        if not nom:
            continue

        site = a_un_site(tags)
        tel = tags.get("phone") or tags.get("contact:phone") or ""
        email = tags.get("email") or tags.get("contact:email") or ""

        if enrichir and site and (not tel or not email):
            print(f"  ... enrichissement : {nom} ({site})")
            e2, t2 = chercher_contact_sur_site(site)
            email = email or e2
            tel = tel or t2

        resultats.append({
            "nom": nom,
            "adresse": extraire_adresse(tags),
            "telephone": tel,
            "email": email,
            "site_web": site,
            "a_un_site": "Non" if not site else "Oui",
            "score_priorite": score_priorite(bool(site), tel, email, nom),
        })

    # dédoublonnage
    vus, dedup = set(), []
    for r in resultats:
        cle = (r["nom"], r["adresse"])
        if cle not in vus:
            vus.add(cle)
            dedup.append(r)

    dedup.sort(key=lambda r: r["score_priorite"], reverse=True)
    return dedup


def ecrire_feuille(ws, lignes):
    colonnes = ["nom", "adresse", "telephone", "email", "site_web", "a_un_site", "score_priorite"]
    ws.append(colonnes)
    for r in lignes:
        ws.append([r[c] for c in colonnes])


def main():
    parser = argparse.ArgumentParser(description="Prospection via OpenStreetMap")
    parser.add_argument("--categorie", help="Catégorie de commerce (voir --liste-categories)")
    parser.add_argument("--toutes", action="store_true", help="Génère toutes les catégories d'un coup")
    parser.add_argument("--ville", help="Ville ou zone, ex: 'Troyes, France'")
    parser.add_argument("--rayon", type=int, default=6000, help="Rayon de recherche en mètres (défaut 6000)")
    parser.add_argument("--no-enrichir", action="store_true", help="Désactive la recherche email/tel sur les sites (plus rapide)")
    parser.add_argument("--liste-categories", action="store_true", help="Affiche les catégories disponibles")
    args = parser.parse_args()

    if args.liste_categories:
        print("Catégories disponibles :")
        for c in sorted(CATEGORIES):
            print(f"  - {c}")
        sys.exit(0)

    if not args.ville or (not args.categorie and not args.toutes):
        parser.error("--ville est requis, avec --categorie OU --toutes")

    if args.categorie and args.categorie not in CATEGORIES:
        print(f"Catégorie inconnue : {args.categorie}")
        print("Utilise --liste-categories pour voir les options.")
        sys.exit(1)

    print(f"Géocodage de '{args.ville}'...")
    lat, lon, nom_complet = geocode_ville(args.ville)
    print(f"  -> {nom_complet} ({lat:.4f}, {lon:.4f})")
    time.sleep(1)

    try:
        from openpyxl import Workbook
    except ImportError:
        print("openpyxl manquant : pip install openpyxl --break-system-packages")
        sys.exit(1)

    enrichir = not args.no_enrichir
    wb = Workbook()
    wb.remove(wb.active)  # on retire la feuille vide par défaut

    categories_a_traiter = list(CATEGORIES) if args.toutes else [args.categorie]
    tous_les_resultats = []

    for i, cat in enumerate(categories_a_traiter):
        lignes = collecter_categorie(cat, lat, lon, args.rayon, enrichir)
        ws = wb.create_sheet(title=cat[:31])  # Excel limite les noms d'onglet à 31 caractères
        ecrire_feuille(ws, lignes)
        for l in lignes:
            l_avec_cat = dict(l)
            l_avec_cat["categorie"] = cat
            tous_les_resultats.append(l_avec_cat)

        if i < len(categories_a_traiter) - 1:
            time.sleep(2)  # pause entre deux catégories, politesse envers Overpass

    if args.toutes:
        tous_les_resultats.sort(key=lambda r: r["score_priorite"], reverse=True)
        ws_tous = wb.create_sheet(title="TOUS", index=0)
        colonnes = ["categorie", "nom", "adresse", "telephone", "email", "site_web", "a_un_site", "score_priorite"]
        ws_tous.append(colonnes)
        for r in tous_les_resultats:
            ws_tous.append([r[c] for c in colonnes])

    if not tous_les_resultats:
        print("Aucun résultat exploitable trouvé. Essaie un rayon plus large.")
        sys.exit(0)

    suffixe = "toutes_categories" if args.toutes else args.categorie
    nom_fichier = f"prospects_{suffixe}_{args.ville.split(',')[0].replace(' ', '_')}.xlsx"
    wb.save(nom_fichier)

    sans_site = sum(1 for r in tous_les_resultats if r["a_un_site"] == "Non")
    avec_email = sum(1 for r in tous_les_resultats if r["email"])
    avec_tel = sum(1 for r in tous_les_resultats if r["telephone"])
    print(f"\n{len(tous_les_resultats)} entreprises trouvées.")
    print(f"  - {sans_site} sans site web")
    print(f"  - {avec_email} avec email trouvé")
    print(f"  - {avec_tel} avec téléphone trouvé")
    print(f"Fichier généré : {nom_fichier}")


if __name__ == "__main__":
    main()
